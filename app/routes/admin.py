from pathlib import Path
from io import BytesIO
from datetime import date

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from flask_sqlalchemy import query
from werkzeug.security import generate_password_hash

from app.extensions import db
from app.models.product import Product
from app.models.quote import Quote
from app.models.sale import Sale
from app.models.user import User
from app.services.sale_service import validate_sale_form
from app.services.upload_service import PAYMENT_PROOF_DIR, save_payment_proof

admin_bp = Blueprint("admin", __name__)


def admin_required(func):
    from functools import wraps

    @wraps(func)
    @login_required
    def wrapper(*args, **kwargs):
        if not current_user.is_authenticated or not current_user.is_admin:
            flash("No tienes permisos para acceder a esta sección.", "error")
            return redirect(url_for("main.dashboard"))
        return func(*args, **kwargs)

    return wrapper


@admin_bp.route("/users", methods=["GET", "POST"])
@admin_required
def users():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "seller")

        if not username or not password:
            flash("Debe completar usuario y contraseña.", "error")
        elif User.query.filter_by(username=username).first():
            flash("El usuario ya existe.", "error")
        else:
            user = User(username=username, password_hash=generate_password_hash(password), role=role)
            db.session.add(user)
            db.session.commit()
            flash("Usuario creado correctamente.", "success")

    users_list = User.query.order_by(User.id.asc()).all()
    return render_template("admin/users.html", users=users_list)


@admin_bp.route("/products", methods=["GET", "POST"])
@admin_required
def products():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        category = request.form.get("category", "general").strip()

        if not name:
            flash("El nombre del producto es obligatorio.", "error")
        else:
            product = Product(name=name, category=category, price=0, is_active=True)
            db.session.add(product)
            db.session.commit()
            flash("Producto agregado correctamente.", "success")

    products_list = Product.query.filter_by(is_active=True).order_by(Product.id.asc()).all()
    return render_template("admin/products.html", products=products_list)


@admin_bp.route("/products/<int:product_id>/delete", methods=["POST"])
@admin_required
def delete_product(product_id):
    product = Product.query.get_or_404(product_id)
    product.is_active = False
    db.session.commit()
    flash("Producto eliminado del catálogo.", "success")
    return redirect(url_for("admin.products"))


@admin_bp.route("/quotes")
@admin_required
def quotes():
    search = request.args.get("q", "").strip()
    status = request.args.get("status", "").strip()

    query = Quote.query

    if search:
        query = query.filter(
            (Quote.quote_number.ilike(f"%{search}%"))
            | (Quote.client_rut.ilike(f"%{search}%"))
            | (Quote.client_razon_social.ilike(f"%{search}%"))
        )

    if status:
        query = query.filter(Quote.status == status)

    quotes_list = query.order_by(Quote.created_at.desc()).all()
    return render_template("admin/quotes.html", quotes=quotes_list, q=search, status=status)


@admin_bp.route("/sales")
@admin_required
def sales():
    seller_id = request.args.get("seller_id", type=int)
    start_raw = (request.args.get("from_date") or "").strip()
    end_raw = (request.args.get("to_date") or "").strip()

    query = Sale.query

    if seller_id:
        query = query.filter(Sale.seller_id == seller_id)

    if start_raw:
        try:
            query = query.filter(Sale.sale_date >= date.fromisoformat(start_raw))
        except ValueError:
            start_raw = ""

    if end_raw:
        try:
            query = query.filter(Sale.sale_date <= date.fromisoformat(end_raw))
        except ValueError:
            end_raw = ""

    sales_list = query.order_by(Sale.sale_date.desc(), Sale.id.desc()).all()
    users_list = User.query.order_by(User.username.asc()).all()

    return render_template(
        "admin/sales.html",
        sales=sales_list,
        users=users_list,
        seller_id=seller_id,
        from_date=start_raw,
        to_date=end_raw,
    )


@admin_bp.route("/sales/<int:sale_id>/export.xlsx")
@admin_required
def export_sale(sale_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sale = Sale.query.get_or_404(sale_id)
    headers = [
        "FECHA", "RUT TITULAR", "NOMBRE Y APELLIDO", "CORREO", "PRODUCTO", "MONTO",
        "RUT FACTURACIÓN", "RAZÓN SOCIAL", "GIRO", "CORREO FACTURACIÓN", "REGIÓN",
        "COMUNA", "DIRECCIÓN", "TELÉFONO", "CORRELATIVO", "TIPO DE VENTA", "MEDIO DE PAGO",
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Ventas"
    worksheet.append(headers)

    worksheet.append([
        sale.sale_date,
        sale.holder_rut,
        sale.holder_name,
        sale.holder_email,
        sale.product,
        float(sale.amount),
        sale.billing_rut,
        sale.billing_legal_name,
        sale.business_activity,
        sale.billing_email,
        sale.region,
        sale.commune,
        sale.address,
        sale.phone,
        sale.invoice_correlative or "",
        sale.sale_type,
        sale.payment_method,
    ])

    header_fill = PatternFill("solid", fgColor="153D32")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["F"].width = 16
    for column_index in range(1, len(headers) + 1):
        column_letter = get_column_letter(column_index)
        if column_letter not in {"A", "F"}:
            worksheet.column_dimensions[column_letter].width = min(
                max(len(headers[column_index - 1]) + 2, 18), 32
            )
    for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=len(headers)):
        row[0].number_format = "dd/mm/yyyy"
        row[5].number_format = '#,##0.00'

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"venta-{sale.id}.xlsx",
    )


@admin_bp.route("/sales/<int:sale_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_sale(sale_id):
    sale = Sale.query.get_or_404(sale_id)
    if request.method == "POST":
        data, errors = validate_sale_form(request.form)
        invoice_correlative = (request.form.get("invoice_correlative") or "").strip()
        if len(invoice_correlative) > 100:
            errors.append("El correlativo no puede superar los 100 caracteres.")

        if errors:
            for error in errors:
                flash(error, "error")
            return render_template("admin/sale_edit.html", sale=sale, form_data=request.form)

        replacement = request.files.get("payment_proof")
        replacement_name = None
        if replacement and replacement.filename:
            replacement_name, error = save_payment_proof(replacement)
            if error:
                errors.append(error)

        if errors:
            for error in errors:
                flash(error, "error")
            return render_template("admin/sale_edit.html", sale=sale, form_data=request.form)

        for field, value in data.items():
            setattr(sale, field, value)
        sale.invoice_correlative = invoice_correlative or None
        if replacement_name:
            previous_name = sale.payment_proof_filename
            sale.payment_proof_filename = replacement_name
            old_path = PAYMENT_PROOF_DIR / Path(previous_name).name
        else:
            old_path = None
        db.session.commit()

        if old_path and old_path.exists():
            old_path.unlink()

        flash("Venta actualizada correctamente.", "success")
        return redirect(url_for("admin.sales"))

    return render_template("admin/sale_edit.html", sale=sale, form_data=None)


@admin_bp.route("/quotes/<int:quote_id>/status", methods=["POST"])
@admin_required
def update_quote_status(quote_id):
    quote = Quote.query.get_or_404(quote_id)
    new_status = (request.form.get("status") or "").strip()

    if new_status not in {"Pendiente", "Completada", "Concretada", "No concretada"}:
        flash("El estado seleccionado no es válido.", "error")
        return redirect(url_for("admin.quotes"))

    quote.status = new_status
    db.session.commit()
    flash("Estado de la cotización actualizado correctamente.", "success")
    return redirect(url_for("admin.quotes"))
