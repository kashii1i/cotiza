import os
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from app import create_app
from app.config import Config
from app.extensions import db
from app.models.product import Product
from app.models.quote import Quote
from app.models.quote_item import QuoteItem
from app.models.sale import Sale
from app.models.user import User


class QuoteWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.pdf_directory = tempfile.TemporaryDirectory()
        self.pdf_directory_patch = patch(
            'app.services.pdf_service.PDF_DIR',
            Path(self.pdf_directory.name),
        )
        self.pdf_directory_patch.start()
        self.upload_directory = tempfile.TemporaryDirectory()
        self.upload_directory_patches = [
            patch('app.services.upload_service.PAYMENT_PROOF_DIR', Path(self.upload_directory.name)),
            patch('app.routes.seller.PAYMENT_PROOF_DIR', Path(self.upload_directory.name)),
            patch('app.routes.admin.PAYMENT_PROOF_DIR', Path(self.upload_directory.name)),
        ]
        for directory_patch in self.upload_directory_patches:
            directory_patch.start()
        with patch.object(Config, 'SQLALCHEMY_DATABASE_URI', 'sqlite:///:memory:'):
            self.app = create_app()
        self.app.config['TESTING'] = True
        self.app_context = self.app.app_context()
        self.app_context.push()
        self.client = self.app.test_client()
        self.client.post('/login', data={'username': 'vendedor', 'password': 'vendedor123'}, follow_redirects=True)

    def tearDown(self):
        db.session.remove()
        self.app_context.pop()
        self.pdf_directory_patch.stop()
        self.pdf_directory.cleanup()
        for directory_patch in reversed(self.upload_directory_patches):
            directory_patch.stop()
        self.upload_directory.cleanup()

    def create_quote(self, seller_username='vendedor'):
        seller = User.query.filter_by(username=seller_username).first()
        quote = Quote(
            quote_number=Quote.generate_number(),
            seller_id=seller.id,
            client_rut='12.345.678-9',
            client_razon_social='Cliente Venta',
            subtotal=10000,
            discount_percent=0,
            discount_amount=0,
            iva=1900,
            total=11900,
        )
        db.session.add(quote)
        db.session.commit()
        return quote

    def sale_form_data(self):
        return {
            'sale_date': '2026-10-02',
            'holder_rut': '12.345.678-9',
            'holder_name': 'Nombre Apellido',
            'holder_email': 'titular@example.cl',
            'product': 'Firma avanzada 1 año',
            'amount': '85000',
            'billing_rut': '76.123.456-7',
            'billing_legal_name': 'Empresa Cliente SpA',
            'business_activity': 'Servicios',
            'billing_email': 'facturacion@example.cl',
            'region': 'Valparaíso',
            'commune': 'Viña del Mar',
            'address': 'Av. Central 123',
            'phone': '+56912345678',
            'sale_type': 'Nueva',
            'payment_method': 'Transferencia',
        }

    def test_invalid_quote_shows_all_errors(self):
        response = self.client.post('/seller/quotes/new', data={
            'client_rut': '',
            'client_razon_social': '',
            'product_id': '',
            'quantity': '',
            'discount_percent': '',
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True).lower()
        self.assertIn('rut', html)
        self.assertIn('razón social', html)
        self.assertIn('producto', html)

    def test_valid_quote_creates_pdf_and_record(self):
        product = Product.query.first()
        response = self.client.post('/seller/quotes/new', data={
            'client_rut': '12.345.678-9',
            'client_razon_social': 'Cliente Demo',
            'product_id': str(product.id),
            'manual_price': '320000',
            'quantity': '2',
            'discount_percent': '10',
            'payment_method': 'Transferencia',
            'due_date': '2027-12-31',
            'bank_details': 'Banco Test',
            'conditions': 'Pago 30 días',
        }, follow_redirects=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Quote.query.count() >= 1, True)

        quote = Quote.query.order_by(Quote.id.desc()).first()
        self.assertIsNotNone(quote)
        self.assertGreater(float(quote.total), 0)
        self.assertEqual(float(quote.discount_amount), 64000)
        self.assertEqual(float(quote.iva), 109440)
        self.assertEqual(float(quote.total), 685440)
        self.assertTrue(os.path.exists(quote.pdf_path))
        with open(quote.pdf_path, 'rb') as generated_pdf:
            self.assertEqual(generated_pdf.read(5), b'%PDF-')

    def test_admin_can_update_quote_status(self):
        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)

        quote = Quote.query.first()
        if quote is None:
            product = Product.query.first()
            self.client.post('/seller/quotes/new', data={
                'client_rut': '1-9',
                'client_razon_social': 'Cliente Admin',
                'product_id': str(product.id),
                'manual_price': '150000',
                'quantity': '1',
                'discount_percent': '0',
            }, follow_redirects=True)
            quote = Quote.query.order_by(Quote.id.desc()).first()

        response = self.client.post(f'/admin/quotes/{quote.id}/status', data={'status': 'Concretada'}, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Quote.query.get(quote.id).status, 'Concretada')

    def test_seller_can_change_quote_status_from_dashboard(self):
        product = Product.query.first()
        self.client.post('/seller/quotes/new', data={
            'client_rut': '45.678.901-2',
            'client_razon_social': 'Cliente Estado',
            'product_id': str(product.id),
            'manual_price': '10000',
            'quantity': '1',
            'discount_percent': '0',
        })
        quote = Quote.query.order_by(Quote.id.desc()).first()

        response = self.client.post(
            f'/seller/quotes/{quote.id}/status',
            data={'status': 'Completada'},
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Quote.query.get(quote.id).status, 'Completada')

    def test_admin_can_remove_products_with_and_without_quote_history(self):
        unused_product = Product(name='Producto para eliminar', category='general', price=1000)
        quoted_product = Product(name='Producto con historial', category='general', price=2000)
        db.session.add_all([unused_product, quoted_product])
        db.session.flush()

        seller = User.query.filter_by(username='vendedor').first()
        quote = Quote(
            quote_number=Quote.generate_number(),
            seller_id=seller.id,
            client_rut='56.789.012-3',
            client_razon_social='Cliente con historial',
            subtotal=2000,
            discount_percent=0,
            discount_amount=0,
            iva=380,
            total=2380,
        )
        db.session.add(quote)
        db.session.flush()
        db.session.add(QuoteItem(
            quote_id=quote.id,
            product_id=quoted_product.id,
            quantity=1,
            unit_price=2000,
            subtotal=2000,
            signer_type='sin_firmador',
        ))
        db.session.commit()

        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)
        self.client.post(f'/admin/products/{unused_product.id}/delete')
        self.client.post(f'/admin/products/{quoted_product.id}/delete')

        self.assertFalse(db.session.get(Product, unused_product.id).is_active)
        self.assertFalse(db.session.get(Product, quoted_product.id).is_active)
        self.assertEqual(QuoteItem.query.filter_by(product_id=quoted_product.id).count(), 1)
        self.assertEqual(QuoteItem.query.filter_by(product_id=quoted_product.id).first().product.name, 'Producto con historial')

    def test_admin_adds_product_without_preset_price(self):
        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)
        self.client.post('/admin/products', data={'name': 'Producto sin precio', 'category': 'general'})

        product = Product.query.filter_by(name='Producto sin precio').first()
        self.assertIsNotNone(product)
        self.assertEqual(float(product.price), 0)

    def test_seller_registers_sale_with_pdf_and_admin_sets_correlative(self):
        response = self.client.post(
            '/seller/sales/new',
            data={**self.sale_form_data(), 'payment_proof': (BytesIO(b'%PDF-1.4 comprobante'), 'comprobante.pdf')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        sale = Sale.query.one()
        self.assertTrue((Path(self.upload_directory.name) / sale.payment_proof_filename).exists())
        self.assertIsNone(sale.invoice_correlative)
        self.assertIsNone(sale.quote_id)

        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)
        admin_response = self.client.post(
            f'/admin/sales/{sale.id}/edit',
            data={**self.sale_form_data(), 'invoice_correlative': 'F-000184'},
            follow_redirects=True,
        )

        self.assertEqual(admin_response.status_code, 200)
        self.assertEqual(db.session.get(Sale, sale.id).invoice_correlative, 'F-000184')
        self.assertIn('Venta actualizada correctamente'.lower(), admin_response.get_data(as_text=True).lower())

    def test_admin_can_open_and_register_a_sale(self):
        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)

        form_response = self.client.get('/seller/sales/new')
        self.assertEqual(form_response.status_code, 200)
        self.assertIn('Ingresar venta', form_response.get_data(as_text=True))

        response = self.client.post(
            '/seller/sales/new',
            data={**self.sale_form_data(), 'payment_proof': (BytesIO(b'%PDF-1.4 comprobante'), 'comprobante.pdf')},
            content_type='multipart/form-data',
        )

        self.assertEqual(response.status_code, 302)
        self.assertIn('/admin/sales', response.headers['Location'])
        sale = Sale.query.one()
        self.assertEqual(sale.seller_id, User.query.filter_by(username='admin').first().id)

    def test_sale_rejects_unsupported_payment_proof_extension(self):
        response = self.client.post(
            '/seller/sales/new',
            data={**self.sale_form_data(), 'payment_proof': (BytesIO(b'not an image'), 'comprobante.exe')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('PDF, PNG, JPG o WEBP', response.get_data(as_text=True))
        self.assertEqual(Sale.query.count(), 0)

    def test_sale_rejects_fake_pdf_content(self):
        response = self.client.post(
            '/seller/sales/new',
            data={**self.sale_form_data(), 'payment_proof': (BytesIO(b'not really a pdf'), 'comprobante.pdf')},
            content_type='multipart/form-data',
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('no coincide con su formato', response.get_data(as_text=True))
        self.assertEqual(Sale.query.count(), 0)

    def test_admin_downloads_one_sale_per_excel_file(self):
        sale = Sale(
            seller_id=User.query.filter_by(username='vendedor').first().id,
            payment_proof_filename='receipt.pdf',
            **{
                **self.sale_form_data(),
                'sale_date': date(2026, 10, 2),
                'amount': 85000,
            },
        )
        other_sale = Sale(
            seller_id=User.query.filter_by(username='vendedor').first().id,
            payment_proof_filename='other-receipt.pdf',
            **{
                **self.sale_form_data(),
                'sale_date': date(2026, 10, 3),
                'holder_name': 'Otra Venta',
                'amount': 42000,
            },
        )
        db.session.add_all([sale, other_sale])
        db.session.commit()
        self.client.post('/logout', follow_redirects=True)
        self.client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=True)

        response = self.client.get(f'/admin/sales/{sale.id}/export.xlsx')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        worksheet = workbook.active
        headers = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        self.assertEqual(len(headers), 17)
        self.assertEqual(headers[0], 'FECHA')
        self.assertEqual(headers[-2:], ('TIPO DE VENTA', 'MEDIO DE PAGO'))
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet['C2'].value, 'Nombre Apellido')
        self.assertEqual(worksheet['F2'].value, 85000)

    def test_fea_and_web_signer_are_quoted_as_separate_products(self):
        fea = Product.query.filter(Product.name.ilike('%firma electrónica avanzada%')).first()
        web_signer = Product.query.filter(Product.name.ilike('%firmador web%')).first()

        valid = self.client.post('/seller/quotes/new', data={
            'client_rut': '22.222.222-2',
            'client_razon_social': 'Cliente FEA con firmador web',
            'product_id': [str(fea.id), str(web_signer.id)],
            'manual_price': ['80000', '30000'],
            'quantity': ['1', '1'],
            'discount_percent': '10',
        }, follow_redirects=True)
        valid_html = valid.get_data(as_text=True)
        self.assertIn('Cliente FEA con firmador web', valid_html)
        quote = Quote.query.order_by(Quote.id.desc()).first()
        self.assertEqual(len(quote.items), 2)
        self.assertEqual({item.product.id for item in quote.items}, {fea.id, web_signer.id})
        self.assertEqual(float(quote.discount_amount), 11000)


if __name__ == '__main__':
    unittest.main()
